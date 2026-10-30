# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import shutil

ROOT = Path(__file__).resolve().parent
FIG_DIR = ROOT / 'figures'
OUT = ROOT / 'Bao_cao_ket_qua_toi_uu_PaySim.docx'
OUT_DOCS = ROOT.parent.parent / 'docs' / 'Bao_cao_ket_qua_toi_uu_PaySim.docx'
OUT_REPO = ROOT.parent.parent / 'Bao_cao_ket_qua_toi_uu_PaySim.docx'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.header_distance = Cm(0.8)
sec.footer_distance = Cm(1)
sec.different_first_page_header_footer = True

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(13)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.line_spacing = 1.06

normal = doc.styles['Normal']
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.first_line_indent = Cm(1)

for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.first_line_indent = Cm(0)
    doc.styles[name].paragraph_format.keep_with_next = True

doc.styles['Heading 1'].paragraph_format.space_before = Pt(5)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(3)
doc.styles['Title'].font.size = Pt(14)
doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hp.paragraph_format.first_line_indent = Cm(0)
hp.paragraph_format.space_after = Pt(0)
r = hp.add_run()
f = OxmlElement('w:fldSimple')
f.set(qn('w:instr'), 'PAGE')
r._r.addnext(f)

def p(text='', bold=False, italic=False, center=False, size=None, space_after=4):
    x = doc.add_paragraph()
    x.paragraph_format.space_after = Pt(space_after)
    if center:
        x.alignment = WD_ALIGN_PARAGRAPH.CENTER
        x.paragraph_format.first_line_indent = Cm(0)
    rr = x.add_run(text)
    rr.bold = bold
    rr.italic = italic
    if size:
        rr.font.size = Pt(size)
    return x

def h(text, level=1):
    return doc.add_paragraph(text, f'Heading {level}')

def page():
    doc.add_page_break()

def add_figure(img_path, width_cm, caption_text):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.first_line_indent = Cm(0)
    pic_p.paragraph_format.space_before = Pt(3)
    pic_p.paragraph_format.space_after = Pt(2)
    pic_p.paragraph_format.keep_with_next = True
    run = pic_p.add_run()
    run.add_picture(str(img_path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(4)
    cr = cap.add_run(caption_text)
    cr.italic = True
    cr.font.size = Pt(11.5)

def cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=12):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    x = cell.paragraphs[0]
    x.paragraph_format.first_line_indent = Cm(0)
    x.paragraph_format.space_before = Pt(2)
    x.paragraph_format.space_after = Pt(2)
    x.paragraph_format.line_spacing = 1.0
    x.alignment = align
    rr = x.add_run(text)
    rr.bold = bold
    rr.font.size = Pt(size)
    return x

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + side)
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '000000')
        borders.append(b)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, val in [('top', '35'), ('bottom', '35'), ('left', '75'), ('right', '75')]:
        z = OxmlElement('w:' + side)
        z.set(qn('w:w'), val)
        z.set(qn('w:type'), 'dxa')
        margins.append(z)
    props.append(margins)
    for i, txt in enumerate(headers):
        t.rows[0].cells[i].width = Cm(widths[i])
        cell_text(t.rows[0].cells[i], txt, True, WD_ALIGN_PARAGRAPH.CENTER)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].width = Cm(widths[i])
            cell_text(cells[i], str(txt), False, WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        z = OxmlElement('w:cantSplit')
        trpr.append(z)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

# Trang 1
t = doc.add_table(rows=1, cols=2)
t.autofit = False
t.columns[0].width = Cm(5.8)
t.columns[1].width = Cm(10.2)
for i, w in enumerate([5.8, 10.2]):
    t.cell(0, i).width = Cm(w)
left, right = t.cell(0, 0), t.cell(0, 1)
cell_text(left, 'NHÓM NGHIÊN CỨU DỰ ÁN\nBIG DATA PAYSIM', True, WD_ALIGN_PARAGRAPH.CENTER, 12)
lp = left.add_paragraph('Số: 01/BC-PAYSIM')
lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
lp.paragraph_format.first_line_indent = Cm(0)
lp.paragraph_format.space_before = Pt(4)
cell_text(right, 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', True, WD_ALIGN_PARAGRAPH.CENTER, 12)
mp = right.add_paragraph()
mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
mp.paragraph_format.first_line_indent = Cm(0)
mp.paragraph_format.space_after = Pt(4)
rr = mp.add_run('Độc lập - Tự do - Hạnh phúc')
rr.bold = True
rr.font.size = Pt(13)
bd = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single')
bottom.set(qn('w:sz'), '4')
bottom.set(qn('w:space'), '3')
bd.append(bottom)
mp._p.get_or_add_pPr().append(bd)
mp.paragraph_format.left_indent = Cm(1.9)
mp.paragraph_format.right_indent = Cm(1.9)
dp = right.add_paragraph()
dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
dp.paragraph_format.first_line_indent = Cm(0)
dp.paragraph_format.space_before = Pt(4)
rr = dp.add_run('Hà Nội, ngày 28 tháng 10 năm 2026')
rr.italic = True
rr.font.size = Pt(13)

doc.add_paragraph('BÁO CÁO', 'Title')
p('Kết quả rà soát toàn diện các hạng mục tối ưu Apache Spark, mô hình học máy\nSpark MLlib và giao diện Dashboard trên hệ thống PaySim Big Data Baseline', bold=True, center=True, size=13.5, space_after=3)
p('Kính gửi: Giảng viên hướng dẫn học phần Dữ liệu lớn (Big Data).', center=True, italic=True, space_after=5)

h('I. ĐẶC ĐIỂM TÌNH HÌNH VÀ KIẾN TRÚC HỆ THỐNG')
p('Hệ thống PaySim Big Data Baseline là nền tảng xử lý và phát hiện gian lận giao dịch tài chính thời gian thực theo kiến trúc Medallion (Bronze – Silver – Gold). Nguồn phát Replayer đọc tập dữ liệu gốc PaySim (6.362.620 giao dịch) đẩy vào Apache Kafka (3 phân vùng); ứng dụng Spark Structured Streaming (Spark 4.0.0, local[2], bộ nhớ 2 GB, chu kỳ micro-batch 10 giây) tiêu thụ luồng sự kiện, lưu trữ bền vững trên HDFS, chấm điểm rủi ro bằng cả tập luật Spark SQL và mô hình học máy Spark MLlib, sau đó ghi kết quả sang MongoDB để hiển thị trực tiếp trên giao diện Streamlit.')
p('Báo cáo tổng hợp kết quả rà soát thực nghiệm đối với 03 trụ cột tối ưu hóa của dự án: (1) Tối ưu luồng xử lý dữ liệu lớn Apache Spark; (2) Cải tiến mô hình học máy phát hiện gian lận Spark MLlib; và (3) Cải tạo giao diện giám sát, khai thác dữ liệu Dashboard. Trong báo cáo, BEFORE chỉ kiến trúc và cấu hình cơ sở ban đầu; AFTER chỉ phiên bản đã cải tiến và đang vận hành thực tế trong mã nguồn hiện hành, đảm bảo giữ nguyên 100% dữ liệu gốc và công thức nghiệp vụ. [1] [2] [3]')
p('Tại thời điểm rà soát ngày 28/10/2026, toàn bộ 7/7 dịch vụ container của hệ thống vận hành ổn định ở trạng thái HEALTHY; khối lưu trữ HDFS đạt 0 block lỗi hoặc thiếu; phiên chạy trực tiếp ghi nhận nguồn phát và Spark xử lý liên tục, đối soát bảo toàn dữ liệu đạt PASS tuyệt đối trên toàn bộ cửa sổ quan sát. [1]')
add_figure(FIG_DIR / 'fig1_architecture.png', 13.6, 'Hình 1. Kiến trúc tổng thể luồng xử lý dữ liệu thời gian thực PaySim Big Data Baseline (Kafka – Spark – HDFS – MLlib – MongoDB – Streamlit).')

# Trang 2
page()
h('II. KẾT QUẢ TỐI ƯU LUỒNG XỬ LÝ APACHE SPARK')
h('1. Tối ưu thu thập thống kê micro-batch (Gom lời gọi thực thi Action)', 2)
p('Nhược điểm của BEFORE: Trong hàm xử lý mỗi micro-batch, thiết kế ban đầu kích hoạt 7 lời gọi thực thi riêng lẻ trên DataFrame (kiểm tra rỗng, đếm tổng bản ghi, tính tổng dung lượng byte, đếm bản ghi hợp lệ, đếm bản ghi trễ hạn và hai lời gọi phân vị approxQuantile riêng cho độ trễ thu nhận và độ trễ đầu cuối). Trong Apache Spark, mỗi action kích hoạt một chu kỳ lập lịch DAG và quét phân vùng độc lập, gây chi phí điều phối lặp lại trên mỗi chu kỳ 10 giây.')
p('Lý do và giải pháp AFTER: Nhằm triệt tiêu các lượt quét dữ liệu dư thừa mà không làm thay đổi công thức tính toán, mã nguồn spark/pipeline.py hợp nhất các phép đo về 3 lời gọi cốt lõi: dùng batch.agg(...) lấy đồng thời số dòng và dung lượng; dùng valid.agg(...) lấy đồng thời số bản ghi hợp lệ và bản ghi trễ; dùng approxQuantile([\"ingest_lag_ms\", \"end_to_end_latency_ms\"], [0.95], 0.05) tính phân vị cho cả hai cột trong một lần quét duy nhất. Giải pháp giúp giảm từ 7 lời gọi xuống còn 3 lời gọi (giảm 57,14% số lời gọi thống kê trong mã nguồn). [1]')

h('2. Tinh chỉnh phân vùng Shuffle tương thích hạ tầng thực thi', 2)
p('Nhược điểm của BEFORE: Cấu hình ban đầu đặt spark.sql.shuffle.partitions = 4 trong khi container Spark vận hành ở chế độ local[2] (chỉ có 2 luồng CPU xử lý song song) với quy mô mỗi micro-batch khoảng 500–700 sự kiện. Việc chia 4 phân vùng shuffle khiến bộ lập lịch phải khởi tạo 4 tác vụ cho mỗi phép tổng hợp theo cửa sổ thời gian và loại giao dịch, làm tăng chi phí tuần tự hóa và chuyển ngữ cảnh CPU.')
p('Lý do và giải pháp AFTER: Để đồng bộ mức độ song song của tầng Shuffle với số lõi thực thi vật lý, hệ thống điều chỉnh spark.sql.shuffle.partitions = 2 (giảm 50% số phân vùng shuffle cấu hình). Kiểm tra trực tiếp qua Spark REST API xác nhận tham số này đang có hiệu lực thực tế, giúp mỗi đợt tổng hợp micro-batch khớp trực tiếp với 2 luồng worker. [1]')

h('3. Tái phân bố dữ liệu trước khi ghi tầng Silver (Repartition theo khóa phân vùng)', 2)
p('Nhược điểm của BEFORE: Khi ghi dữ liệu sạch xuống HDFS tầng Silver theo cấu trúc thư mục phân vùng partitionBy(\"run_id\", \"type\"), phương án BEFORE ghi trực tiếp từ các phân vùng bộ nhớ hiện hữu. Do các bản ghi cùng loại giao dịch (PAYMENT, TRANSFER, CASH_OUT, CASH_IN, DEBIT) nằm rải rác trên nhiều task, mỗi task đều mở luồng ghi riêng vào cùng một thư mục đích, làm phát sinh nhiều tệp Parquet nhỏ (Small Files Problem) và tăng áp lực quản lý metadata lên HDFS NameNode.')
p('Lý do và giải pháp AFTER: Phiên bản AFTER bổ sung bước .repartition(\"run_id\", \"type\") ngay trước lệnh ghi Silver, gom toàn bộ bản ghi có cùng khóa phân vùng về chung một phân vùng xử lý trước khi ghi xuống đĩa, qua đó giảm số trình ghi đồng thời trên mỗi thư mục loại giao dịch. [1]')

h('4. Kết quả vận hành thực tế sau tối ưu Spark', 2)
table(['Chỉ tiêu vận hành hệ thống Spark – HDFS', 'Giá trị ghi nhận thực tế'], [
    ('Quy mô cửa sổ quan sát (60 micro-batch từ 981 đến 1040)', '40.200 sự kiện giao dịch'),
    ('Thời gian xử lý trung bình 01 micro-batch', '11,239 giây (nhỏ nhất 5,321 giây)'),
    ('Thông lượng xử lý gộp thực tế (Processing Throughput)', '59,62 sự kiện/giây (> 50 sự kiện/giây đầu vào)'),
    ('Trạng thái đối soát bảo toàn dữ liệu (Data Reconciliation)', 'PASS tuyệt đối (60/60 micro-batch, độ lệch = 0)'),
    ('Sức khỏe khối lưu trữ HDFS NameNode & DataNode', 'HEALTHY; 0 corrupt blocks; 0 missing blocks'),
], [9.2, 6.8])
p('Nhận xét: Luồng Spark Structured Streaming sau tối ưu đạt thông lượng xử lý trung bình 59,62 sự kiện/giây, vượt tốc độ bơm dữ liệu của nguồn phát (50,0 sự kiện/giây), giúp hệ thống hấp thụ hết hàng đợi Kafka và duy trì trạng thái đối soát nhất quán tuyệt đối. [1]')

# Trang 3
page()
h('III. KẾT QUẢ CẢI TIẾN MÔ HÌNH HỌC MÁY SPARK MLLIB')
h('1. Nhược điểm của mô hình BEFORE (Logistic Regression) và lý do cải tiến', 2)
p('Nhược điểm của BEFORE: Mô hình cơ sở sử dụng thuật toán Hồi quy Logistic (LogisticRegression) trên không gian đặc trưng gồm loại giao dịch (One-Hot Encoding), logarit số tiền log1p(amount) và khung giờ trong ngày. Do biên quyết định tuyến tính đối với log-odds, Logistic Regression không thể tự biểu diễn các tương tác phi tuyến bậc cao (ví dụ: gian lận chỉ tập trung ở giao dịch TRANSFER hoặc CASH_OUT khi giá trị giao dịch vượt ngưỡng lớn trong một số khung giờ nhất định). Trên tập dữ liệu PaySim mất cân bằng cực단 (tỷ lệ gian lận ~0,13%), việc gắn trọng số lớp lớn đẩy siêu phẳng tuyến tính lệch mạnh, gây ra tới 14.522 cảnh báo giả trên tập kiểm thử nhưng chỉ bắt được 43/123 giao dịch gian lận (Recall chỉ đạt 34,96%). [2]')
p('Lý do cải tiến: Nhằm nắm bắt trực tiếp các quy tắc rẽ nhánh phi tuyến giữa loại giao dịch, quy mô tiền và thời điểm phát sinh mà không làm phình chiều vector đặc trưng, dự án thực nghiệm đối chứng 4 thuật toán trong Spark MLlib trên cùng tập đặc trưng, cùng trọng số lớp và ngưỡng phân loại 0,50, lấy chỉ số Average Precision (AP – diện tích dưới đường cong Precision-Recall) trên tập thẩm định làm tiêu chí lựa chọn mô hình. [2]')

h('2. Giải pháp AFTER và kết quả thực nghiệm đối chứng trên tập kiểm thử', 2)
p('Giải pháp AFTER lựa chọn Cây quyết định (DecisionTreeClassifier, giới hạn độ sâu maxDepth = 5, số mẫu tối thiểu mỗi nút lá minInstancesPerNode = 20) đạt điểm AP trên tập Validation cao nhất (14,525%, vượt xa Logistic Regression 0,113%). Kết quả kiểm định độc lập trên tập Test (246.565 giao dịch, chứa 123 giao dịch gian lận ở các bước thời gian 33–40) như sau:')
table(['Mô hình thực nghiệm', 'AP (%)', 'Recall (%)', 'F1 (%)', 'Bắt đúng (TP)', 'Báo động giả (FP)'], [
    ('Logistic Regression (BEFORE)', '0,250', '34,96', '0,586', '43 / 123', '14.522'),
    ('Decision Tree (AFTER - Đang dùng)', '11,698', '50,41', '0,981', '62 / 123', '12.454'),
    ('Random Forest (Đối chứng)', '1,082', '52,03', '1,237', '64 / 123', '10.158'),
    ('GBTClassifier (Đối chứng)', '4,895', '28,46', '0,987', '35 / 123', '6.933'),
], [5.1, 1.9, 2.2, 1.8, 2.4, 2.6])
p('Hiệu quả cải tiến định lượng: So với mô hình BEFORE, mô hình AFTER (Decision Tree) tăng 46,79 lần chỉ số Average Precision (từ 0,250% lên 11,698%); tăng Recall thêm 15,45 điểm phần trăm (bắt thêm 19 giao dịch gian lận, đạt 62/123); đồng thời giảm 2.068 cảnh báo giả (giảm 14,24% FP) và tăng 67,41% điểm F1. Trên tập kiểm toán độc lập (Audit, bước 41–48 gồm 108 ca gian lận), trong khi Logistic Regression bỏ sót hoàn toàn (0/108), Decision Tree vẫn phát hiện thành công 44/108 ca gian lận (Recall đạt 40,74%). [2]')
add_figure(FIG_DIR / 'fig2_mllib_comparison.png', 14.8, 'Hình 2. Đối chiếu trực quan hiệu quả phát hiện giữa tập luật SQL Rule và mô hình Spark MLlib cùng Ma trận nhầm lẫn trên luồng đánh giá.')

h('3. Đánh giá hiệu năng suy luận và giới hạn vận hành', 2)
p('Thời gian chấm điểm (scoring latency kèm ghi Parquet Gold) của Decision Tree đạt 29,88 ms / 1.000 dòng, tương đương tốc độ của Logistic Regression (26,04 – 33,47 ms / 1.000 dòng), đảm bảo không gây nghẽn cổ chai cho luồng Streaming. Hạn chế còn tồn tại là tỷ lệ Precision tuyệt đối ở ngưỡng 0,50 vẫn thấp (~0,495%) do dữ liệu gốc chỉ sử dụng 3 đặc trưng giao dịch tĩnh, đòi hỏi kết hợp hậu kiểm nghiệp vụ. [2]')

# Trang 4
page()
h('IV. KẾT QUẢ CẢI TIẾN GIAO DIỆN GIÁM SÁT DASHBOARD')
h('1. Nhược điểm của giao diện BEFORE và yêu cầu cải tiến', 2)
p('Nhược điểm của BEFORE: Giao diện ban đầu (PaySim — Live Big Data Control Room) bộc lộ 4 hạn chế về khả năng khai thác: (1) Bảng giao dịch mới nhất chỉ hiển thị danh sách tĩnh, hoàn toàn thiếu công cụ lọc kết hợp, thiếu ô tìm kiếm theo mã sự kiện (Event ID) và không hỗ trợ tải xuống CSV; (2) Khi giao diện tự động làm mới sau mỗi 2 giây, thứ tự bảng dịch chuyển liên tục khiến người vận hành mất dấu bản ghi đang kiểm tra; (3) Khu vực 10 nhóm chỉ số (10 optimization tracks) dàn trải trên một mặt phẳng dài gây quá tải thị giác; (4) Khối đồ họa Hero dễ bị cắt xén đường viền hoặc khởi động lại hiệu ứng sau mỗi chu kỳ làm mới. [3]')

h('2. Giải pháp cải tiến AFTER đã triển khai trên hệ thống', 2)
p('Thứ nhất, tái cấu trúc không gian làm việc thành 4 phân hệ chuyên biệt trên thanh điều hướng có biểu tượng trực quan (📡 Live Control Room, 🎯 Metrics theo 10 tracks, 🛡️ Fraud & Business, 🔎 New data streams), kết hợp bảng màu Pastel đồng nhất, khung ruy-băng kính 3D (468px) không bị cắt xén và thẻ trạng thái Read-only workspace căn chỉnh thẳng hàng dưới thẻ phân tích. [3]')
p('Thứ hai, bổ sung bộ công cụ khai thác giao dịch đa chiều (Transaction Explorer): cho phép lọc kết hợp theo loại giao dịch, trạng thái cảnh báo, khoảng số tiền tối thiểu/tối đa, thanh trượt khoảng điểm rủi ro (Risk score 0,00 – 1,00), tìm kiếm chuỗi con Event ID, ghim bản ghi đang xem qua các chu kỳ làm mới 2 giây và xuất tệp CSV đúng tập kết quả đang lọc. [3]')
p('Thứ ba, cô lập chu kỳ làm mới (st.fragment(run_every=\"2s\")) tách biệt khỏi phần khung tĩnh Hero và áp dụng cơ chế tạm dừng chuyển động khi ẩn tab trình duyệt hoặc bật chế độ giảm chuyển động (prefers-reduced-motion), giúp giao diện vận hành mượt mà và bảo toàn 100% công thức tính toán KPI. [3]')

add_figure(FIG_DIR / 'fig3_dashboard_before_after.png', 15.2, 'Hình 3. Đối chiếu giao diện tổng quan BEFORE (Control Room cơ sở, trái) và AFTER (PaySim Data Workspace, phải) bảo toàn 100% số liệu KPI.')
add_figure(FIG_DIR / 'fig4_transactions_before_after.png', 15.2, 'Hình 4. So sánh phân hệ tra cứu giao dịch BEFORE (bảng tĩnh không có bộ lọc, trái) và AFTER (bộ lọc đa tiêu chí, tìm kiếm Event ID & Risk Score, phải).')

# Trang 5
page()
h('3. Kết quả kiểm thử tự động và đối soát bảo toàn dữ liệu Dashboard', 2)
table(['Hạng mục kiểm định chất lượng Dashboard', 'Kết quả nghiệm thu thực tế'], [
    ('Bộ kiểm thử đơn vị và tích hợp tự động (Unit & Integration Tests)', '34 bài đạt (PASS); 0 bài lỗi; 12 subtests đạt tuyệt đối'),
    ('Kiểm thử phát lại giao diện trên trình duyệt (Playwright Replay)', '12/12 ảnh chụp chuẩn hóa (8 Desktop 1600x1100; 4 Mobile 390x844)'),
    ('Đối soát bảo toàn số liệu nghiệp vụ giữa BEFORE và AFTER', 'Khớp chính xác 100% cả 6/6 chỉ số KPI trên cả 4 phân hệ'),
    ('Kiểm tra lỗi ngoại lệ runtime và tràn bố cục ngang (Overflow)', '0 ngoại lệ Streamlit; 0 lỗi tràn khung ngang ở mọi độ phân giải'),
    ('Kiểm chứng tính năng lọc, ghim bản ghi và xuất dữ liệu CSV', 'Lọc chính xác 60->8 bản ghi; giữ đúng bản ghi ghim; xuất CSV khớp 100%'),
], [8.8, 7.2])
p('Nguồn đối chiếu: Hồ sơ kiểm thử tự động docs/evidence/dashboard-english và dashboard-motion. Mọi thao tác lọc và tra cứu trên Dashboard chỉ tác động lên lớp trình bày trong bộ nhớ đệm (60 micro-batch, 60 giao dịch gần nhất), tuyệt đối không làm biến đổi dữ liệu gốc hay các chỉ số KPI toàn cục của lượt chạy. [3]')

h('V. ĐÁNH GIÁ CHUNG VÀ KIẾN NGHỊ')
h('1. Đánh giá chung', 2)
p('Qua rà soát thực nghiệm toàn diện sau khi vận hành hệ thống, cả 03 hạng mục tối ưu đều hoàn thành đúng mục tiêu kỹ thuật đề ra mà không vi phạm nguyên tắc bảo toàn dữ liệu gốc: (1) Tầng Apache Spark giảm 57,14% số lời gọi thu thập thống kê trong mỗi micro-batch, đồng bộ phân vùng shuffle (2 phân vùng) với số luồng thực thi và gom nhóm dữ liệu theo khóa trước khi ghi Silver, duy trì thông lượng 59,62 sự kiện/giây vượt tốc độ nguồn phát; (2) Tầng Spark MLlib thay thế Hồi quy Logistic bằng Cây quyết định (Decision Tree), giúp tăng 46,79 lần chỉ số Average Precision, nâng tỷ lệ phát hiện gian lận từ 34,96% lên 50,41% và giảm 14,24% cảnh báo giả trên tập kiểm thử; (3) Tầng Dashboard chuyển đổi từ màn hình theo dõi tĩnh sang không gian làm việc tương tác có bộ lọc đa tiêu chí, ghim bản ghi qua chu kỳ làm mới 2 giây và trích xuất CSV, vượt qua toàn bộ 34 kịch bản kiểm thử tự động. [1] [2] [3]')

h('2. Kiến nghị hướng phát triển tiếp theo', 2)
p('Một là, đối với tầng xử lý Apache Spark: Thiết lập quy trình đo kiểm A/B cố định hạt giống dữ liệu và tách biệt giai đoạn khởi động (warm-up) nhằm định lượng chính xác tỷ lệ giảm số tệp Parquet nhỏ trên HDFS và mức cải thiện độ trễ trung vị của từng thay đổi cấu hình.')
p('Hai là, đối với tầng mô hình học máy Spark MLlib: Bổ sung các đặc trưng lịch sử hành vi theo tài khoản (tần suất giao dịch trượt theo cửa sổ thời gian, tỷ lệ rút cạn số dư oldbalanceOrg -> newbalanceOrig) và tinh chỉnh ngưỡng phân loại theo chi phí nghiệp vụ nhằm nâng cao độ chính xác (Precision), giảm tải cho khâu hậu kiểm.')
p('Ba là, đối với tầng Dashboard: Mở rộng cơ chế truy vấn phân trang phía máy chủ (server-side pagination & filtering) trực tiếp trên tập hợp MongoDB/HDFS Gold khi cần tra cứu lịch sử giao dịch vượt quá cửa sổ 60 bản ghi thời gian thực.')

h('3. Tài liệu và hồ sơ chứng cứ đối chiếu', 2)
refs = [
    '[1] Mã nguồn spark/pipeline.py; dữ liệu vận hành MongoDB, Spark REST API và HDFS ngày 28/10/2026; hồ sơ artifacts/report-20261028/evidence/spark-review.md.',
    '[2] Tài liệu docs/04_Huong_Dan_Doc_Va_Toi_Uu_MLlib.md (Mục 8–10); hồ sơ thực nghiệm mô hình artifacts/report-20261028/evidence/model-review.json.',
    '[3] Hồ sơ kiểm định giao diện docs/evidence/dashboard-english/verification.md, browser-results.json và docs/evidence/dashboard-motion/results.json.',
    '[4] Thể thức văn bản trình bày theo Phụ lục I, Nghị định số 30/2020/NĐ-CP ngày 05/3/2020 của Chính phủ về công tác văn thư.',
]
for txt in refs:
    x = p(txt, size=12, space_after=2)
    x.paragraph_format.first_line_indent = Cm(0)
    x.paragraph_format.line_spacing = 1.0

t = doc.add_table(rows=1, cols=2)
t.autofit = False
t.columns[0].width = Cm(8)
t.columns[1].width = Cm(8)
cell_text(t.cell(0, 0), 'Nơi nhận:\n- Như kính gửi;\n- Lưu: Hồ sơ dự án PaySim.', False, WD_ALIGN_PARAGRAPH.LEFT, 11.5)
cell_text(t.cell(0, 1), 'TM. NHÓM THỰC HIỆN DỰ ÁN\nNGƯỜI LẬP BÁO CÁO\n(Ký và ghi rõ họ tên)\n\n\n', True, WD_ALIGN_PARAGRAPH.CENTER, 12)

doc.core_properties.title = 'Báo cáo kết quả tối ưu dự án PaySim Big Data Baseline'
doc.core_properties.subject = 'Báo cáo kết quả tối ưu Apache Spark, Spark MLlib và Dashboard'
doc.core_properties.author = 'Nhóm nghiên cứu dự án PaySim'
doc.core_properties.keywords = 'PaySim, Apache Spark, Spark MLlib, Dashboard, Nghị định 30/2020/NĐ-CP'
doc.save(OUT)
shutil.copyfile(OUT, OUT_DOCS)
shutil.copyfile(OUT, OUT_REPO)
print(OUT)
